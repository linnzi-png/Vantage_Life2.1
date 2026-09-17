"""Morgans Secret Sauce — the one-sheet weekly recognition report.

Per owner (2026-09-17): a deliberately simple workbook, one sheet, four blocks,
covering every agent and leader across all offices for one reporting week.
Nothing here re-imports and nothing computes a rate; it ranks.

    Top 4 Producers        veterans, by Gross ALP
    Top Rookie Producers   rookies,  by Gross ALP
    Top Plus Lead Sales    everyone, by Ref Sales      ("Plus Lead" = referral)
    Top Plus Leads Collected  everyone, by Referrals obtained

Each block is three columns — MGA / Agent / value — laid out two blocks across,
matching the sheet the office already builds by hand. Leaders are ranked
alongside agents on their own production; the MGA column is the person's
upline MGA (blank for an MGA themselves and anyone above).

Rookie vs veteran comes off agent_profiles.is_rookie. Anyone not flagged as a
rookie ranks as a veteran, so the sheet still fills while tenures are being
set from the roster.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TOP_N = 4

NAVY = "1F2D6B"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF", size=12)
SUB_FONT = Font(bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(size=11)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (title, value header, metric key, filter, number format)
#   filter: "veteran" | "rookie" | None (everyone)
Block = Tuple[str, str, str, Optional[str], str]
BLOCKS: List[Block] = [
    ("Top 4 Producers", "ALP", "gross_alp", "veteran", '"$"#,##0'),
    ("Top Rookie Producers", "ALP", "gross_alp", "rookie", '"$"#,##0'),
    ("Top Plus Lead Sales", "Sales", "ref_sales", None, "0"),
    ("Top Plus Leads Collected", "Refs", "refs_obtained", None, "0"),
]

# Where each block starts: (row, column) — two across, two down, with a gap
# column D between them and a blank row 7 between the pairs.
_ANCHORS = [(1, 1), (1, 5), (8, 1), (8, 5)]


def rank(rows: List[Dict[str, Any]], metric: str, tenure: Optional[str]) -> List[Dict[str, Any]]:
    """Top N by `metric`, descending, ties broken by name. Zero producers are
    never ranked — an empty slot is more honest than a name with a 0 beside it."""
    pool = rows
    if tenure == "rookie":
        pool = [r for r in rows if r.get("is_rookie") is True]
    elif tenure == "veteran":
        pool = [r for r in rows if r.get("is_rookie") is not True]
    pool = [r for r in pool if float(r.get(metric) or 0) > 0]
    pool.sort(key=lambda r: (-float(r.get(metric) or 0), (r.get("name") or "").lower()))
    return pool[:TOP_N]


def _write_block(ws, row: int, col: int, block: Block, rows: List[Dict[str, Any]]) -> None:
    title, value_head, metric, tenure, fmt = block
    c1, c3 = get_column_letter(col), get_column_letter(col + 2)

    ws.merge_cells(f"{c1}{row}:{c3}{row}")
    t = ws.cell(row, col, title)
    t.fill, t.font = HEAD_FILL, HEAD_FONT
    t.alignment = Alignment(horizontal="center")
    for i in range(3):
        ws.cell(row, col + i).fill = HEAD_FILL

    for i, head in enumerate(("MGA", "Agent", value_head)):
        h = ws.cell(row + 1, col + i, head)
        h.fill, h.font, h.border = HEAD_FILL, SUB_FONT, BOX

    ranked = rank(rows, metric, tenure)
    for n in range(TOP_N):
        r = row + 2 + n
        person = ranked[n] if n < len(ranked) else None
        cells = [
            ws.cell(r, col, (person.get("mga") or "").upper() if person else None),
            ws.cell(r, col + 1, person.get("name") if person else None),
            ws.cell(r, col + 2, person.get(metric) if person else None),
        ]
        for c in cells:
            c.font, c.border = BODY_FONT, BOX
        cells[2].number_format = fmt
        cells[2].alignment = Alignment(horizontal="right")


def build_workbook(week_start: date, rows: List[Dict[str, Any]]) -> io.BytesIO:
    """`rows`: one per person — name, mga (upline MGA name or None), is_rookie,
    gross_alp, ref_sales, refs_obtained — summed over the week."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Week of {week_start.month}.{week_start.day}.{week_start.year % 100}"

    for block, (row, col) in zip(BLOCKS, _ANCHORS):
        _write_block(ws, row, col, block, rows)

    for letter, width in (("A", 24), ("B", 22), ("C", 11), ("D", 4),
                          ("E", 24), ("F", 22), ("G", 11)):
        ws.column_dimensions[letter].width = width
    ws.sheet_view.showGridLines = False

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
