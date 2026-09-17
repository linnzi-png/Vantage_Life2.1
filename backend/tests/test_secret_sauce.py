"""Morgans Secret Sauce — the one-sheet weekly recognition report."""
import io

import openpyxl

import secret_sauce
from conftest import auth, make_session

WEEK = "2026-07-01"  # a Wednesday


async def _add(db, agent_id, office, sales_day, **metrics):
    doc = {
        "entry_id": f"pe_{agent_id}_{sales_day}", "agent_id": agent_id, "office": office,
        "sales_day": sales_day, "sets": 0, "sits": 0, "sales": 0, "ots_sits": 0, "ots_sales": 0,
        "n1": 0, "refs_obtained": 0, "ref_sits": 0, "ref_sales": 0, "pos_sits": 0, "pos_sales": 0,
        "vet_sits": 0, "vet_sales": 0, "gross_alp": 0.0, "net_alp": 0.0,
    }
    doc.update(metrics)
    await db.production_entries.insert_one(doc)


async def _rga(db):
    return await make_session(db, role="level_4", agent_id="RGA_1", email="rga1@test.dev")


async def _sheet(client, db):
    r = await client.get(f"/api/vault/secret-sauce?week_start={WEEK}", headers=auth(await _rga(db)))
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert "Morgans_Secret_Sauce" in r.headers["content-disposition"]
    return openpyxl.load_workbook(io.BytesIO(r.content)).active


def _anchor(ws, title):
    """(row, col) of the block with this title — blocks move down as ties grow."""
    for row in ws.iter_rows():
        for c in row:
            if c.value == title:
                return c.row, c.column
    raise AssertionError(f"no block titled {title!r}")


def _block(ws, title, n=secret_sauce.TOP_N):
    """The first `n` data rows under the block: [(mga, name, value)]."""
    row, col = _anchor(ws, title)
    return [(ws.cell(row + 2 + i, col).value, ws.cell(row + 2 + i, col + 1).value,
             ws.cell(row + 2 + i, col + 2).value) for i in range(n)]


VETS, ROOKIES = "Top 4 Producers", "Top Rookie Producers"
PLUS_SALES, PLUS_REFS = "Top Plus Lead Sales", "Top Plus Leads Collected"


# ---------------- layout ----------------

async def test_four_blocks_two_across(client, seeded_db):
    ws = await _sheet(client, seeded_db)
    assert _anchor(ws, VETS) == (1, 1)
    assert _anchor(ws, ROOKIES) == (1, 5)
    assert _anchor(ws, PLUS_SALES) == (8, 1), "4 rows + a blank row below the upper pair"
    assert _anchor(ws, PLUS_REFS) == (8, 5)
    assert (ws["A2"].value, ws["B2"].value, ws["C2"].value) == ("MGA", "Agent", "ALP")
    assert (ws["E9"].value, ws["F9"].value, ws["G9"].value) == ("MGA", "Agent", "Refs")
    assert ws["C9"].value == "Sales"


# ---------------- ranking ----------------

async def test_veterans_and_rookies_rank_separately_on_weekly_alp(client, seeded_db):
    await seeded_db.agent_profiles.update_one({"agent_id": "AG_2"}, {"$set": {"is_rookie": True}})
    await _add(seeded_db, "AG_1", "MCM", "2026-07-01", gross_alp=1000.0)
    await _add(seeded_db, "AG_1", "MCM", "2026-07-03", gross_alp=500.0)   # same week: summed
    await _add(seeded_db, "GA_1", "MCM", "2026-07-02", gross_alp=2000.0)  # a leader ranks too
    await _add(seeded_db, "AG_2", "AMP", "2026-07-02", gross_alp=300.0)   # rookie
    await _add(seeded_db, "AG_1", "MCM", "2026-07-08", gross_alp=9999.0)  # next week: excluded
    ws = await _sheet(client, seeded_db)

    vets = _block(ws, VETS)
    assert vets[0] == ("MGA ONE", "Ga One", 2000.0)
    assert vets[1] == ("MGA ONE", "Agent One", 1500.0)
    assert vets[2] == (None, None, None), "no zero producers, no padding names"

    rookies = _block(ws, ROOKIES)
    assert rookies[0] == ("MGA ONE", "Agent Two", 300.0)
    assert rookies[1] == (None, None, None)
    assert "Agent Two" not in [v[1] for v in vets], "a rookie never ranks as a veteran"


async def test_plus_lead_blocks_use_referral_metrics(client, seeded_db):
    await _add(seeded_db, "AG_1", "MCM", "2026-07-01", refs_obtained=45, ref_sales=1, sales=6)
    await _add(seeded_db, "AG_2", "AMP", "2026-07-01", refs_obtained=26, ref_sales=3, sales=1)
    ws = await _sheet(client, seeded_db)
    # Plus Lead Sales = ref_sales, NOT total sales.
    assert [r[1:] for r in _block(ws, PLUS_SALES)[:2]] == [("Agent Two", 3), ("Agent One", 1)]
    # Plus Leads Collected = refs_obtained.
    assert [r[1:] for r in _block(ws, PLUS_REFS)[:2]] == [("Agent One", 45), ("Agent Two", 26)]


async def test_top_n_is_top_n_values(client, seeded_db):
    for i, aid in enumerate(["AG_1", "AG_2", "GA_1", "GA_2", "SA_1", "MGA_1"]):
        await _add(seeded_db, aid, "MCM", "2026-07-01", gross_alp=100.0 * (i + 1))
    ws = await _sheet(client, seeded_db)
    vets = _block(ws, VETS)
    assert [v[2] for v in vets] == [600.0, 500.0, 400.0, 300.0]
    # An MGA has no MGA above them: the column is blank, not their own name.
    assert vets[0][1] == "Mga One" and vets[0][0] is None
    assert ws.cell(7, 1).value is None, "six distinct values: only four make the list"


async def test_ties_are_all_listed_and_push_the_lower_blocks_down(client, seeded_db):
    """Owner's example, on a top-3: 16, 12, 12, 10 lists four names. Here on
    the top-4 refs block: 16, 12, 12, 10, 8, 8 lists six, and the sheet grows
    rather than cutting anyone."""
    for aid, refs in [("AG_1", 16), ("AG_2", 12), ("GA_1", 12), ("GA_2", 10),
                      ("SA_1", 8), ("MGA_1", 8), ("RGA_1", 5)]:
        await _add(seeded_db, aid, "MCM", "2026-07-01", refs_obtained=refs)
    ws = await _sheet(client, seeded_db)
    refs = _block(ws, PLUS_REFS, n=7)
    assert [r[2] for r in refs] == [16, 12, 12, 10, 8, 8, None]
    assert [r[1] for r in refs[1:3]] == ["Agent Two", "Ga One"], "A-Z inside a tie"
    assert "Rga One" not in [r[1] for r in refs], "5 is the fifth value: out"

    # Upper pair is empty (4 rows), lower pair still lands on row 8 — and the
    # lower pair's own height is whatever the taller block needs.
    assert _anchor(ws, PLUS_REFS) == (8, 5)
    assert ws.cell(8 + 2 + 5, 6).value == "Sa One"  # sixth data row exists (Mga < Sa A-Z)


async def test_a_tie_in_the_upper_pair_moves_the_lower_pair(client, seeded_db):
    for i, aid in enumerate(["AG_1", "AG_2", "GA_1", "GA_2", "SA_1"]):
        await _add(seeded_db, aid, "MCM", "2026-07-01", gross_alp=100.0 if i else 200.0)
    ws = await _sheet(client, seeded_db)
    # 200, then four people tied on 100: five rows in the veterans block.
    assert [v[2] for v in _block(ws, VETS, n=5)] == [200.0, 100.0, 100.0, 100.0, 100.0]
    assert _anchor(ws, PLUS_SALES) == (9, 1), "one row lower than the default layout"
    assert _anchor(ws, PLUS_REFS) == (9, 5)


# ---------------- access ----------------

async def test_finance_admin_and_admin_can_pull_it(client, seeded_db):
    await seeded_db.agent_profiles.insert_one({
        "agent_id": "FA_1", "name": "Finance One", "email": "fa1@test.dev",
        "role": "finance_admin", "upline_id": None, "office": "",
    })
    fa = await make_session(seeded_db, role="finance_admin", agent_id="FA_1", email="fa1@test.dev")
    assert (await client.get(f"/api/vault/secret-sauce?week_start={WEEK}", headers=auth(fa))).status_code == 200


async def test_mga_is_refused(client, seeded_db):
    token = await make_session(seeded_db, role="level_3", agent_id="MGA_1", email="mga1@test.dev")
    assert (await client.get(f"/api/vault/secret-sauce?week_start={WEEK}", headers=auth(token))).status_code == 403


async def test_week_start_must_be_a_wednesday(client, seeded_db):
    r = await client.get("/api/vault/secret-sauce?week_start=2026-07-02", headers=auth(await _rga(seeded_db)))
    assert r.status_code == 400
