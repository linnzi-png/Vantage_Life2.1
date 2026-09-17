"""Code-date import: rookie = coded within the last 12 months, rolling."""
import io
from datetime import date, datetime

import openpyxl

import code_dates
from conftest import auth, make_session


def _workbook(rows, header=("Agent", "Code Date", "MGA Name"), extra_sheet=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Producers"
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    if extra_sheet:
        title, hdr, xrows = extra_sheet
        ws2 = wb.create_sheet(title)
        ws2.append(list(hdr))
        for r in xrows:
            ws2.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _admin(db):
    tok = await make_session(db, role="level_4", agent_id="RGA_1", email="rga1@test.dev")
    await db.users.update_one({"email": "rga1@test.dev"}, {"$set": {"is_admin": True}})
    return tok


async def _upload(client, token, content, dry_run=False):
    return await client.post(
        "/api/admin/import-code-dates", headers=auth(token),
        files={"file": ("RGA Code Dates.xlsx", content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"dry_run": "true" if dry_run else "false"},
    )


# ---------------- the rule ----------------

def test_rookie_window_is_twelve_months_rolling():
    on = date(2026, 9, 17)
    assert code_dates.rookie_on("2025-09-18", on) is True     # inside the year
    assert code_dates.rookie_on("2025-09-17", on) is False    # anniversary: veteran
    assert code_dates.rookie_on("2016-08-23", on) is False
    assert code_dates.rookie_on(None, on) is None             # nothing to judge by


# ---------------- the parser ----------------

def test_parser_reads_every_sheet_by_header_not_position():
    content = _workbook(
        [("AGENT ONE", datetime(2026, 3, 2), "MGA ONE"), ("", None, None)],
        extra_sheet=("MGA", ("Leader", "Associate ID", "Highest Contract", "MGA", "Code Date"),
                     [("Mark Neilson", 1, "Master General Agent", "MGA ONE", datetime(2022, 7, 1))]),
    )
    rows = code_dates.parse_workbook(io.BytesIO(content))
    assert [(r["sheet"], r["name"], r["code_date"]) for r in rows] == [
        ("All Producers", "AGENT ONE", "2026-03-02"),
        ("MGA", "MGA ONE", "2022-07-01"),   # the MGA column, never "Leader"
    ]


# ---------------- the import ----------------

async def test_import_sets_code_date_and_rookie_flag(client, seeded_db):
    token = await _admin(seeded_db)
    content = _workbook([
        ("AGENT ONE", datetime(2026, 3, 2), "MGA ONE"),        # rookie
        ("Agent Two", datetime(2016, 8, 23), "MGA ONE"),       # veteran
        ("Qaradaghi, Snoor", datetime(2024, 2, 22), "X"),      # not on roster
    ])
    r = await _upload(client, token, content, dry_run=True)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["dry_run"] is True
    assert (body["matched"], body["rookies"], body["veterans"]) == (2, 1, 1)
    assert body["unmatched"] == ["Qaradaghi, Snoor"]
    # Dry run wrote nothing.
    assert "code_date" not in await seeded_db.agent_profiles.find_one({"agent_id": "AG_1"})

    r = await _upload(client, token, content)
    assert r.status_code == 200, r.text
    a1 = await seeded_db.agent_profiles.find_one({"agent_id": "AG_1"})
    a2 = await seeded_db.agent_profiles.find_one({"agent_id": "AG_2"})
    assert (a1["code_date"], a1["is_rookie"]) == ("2026-03-02", True)
    assert (a2["code_date"], a2["is_rookie"]) == ("2016-08-23", False)
    assert await seeded_db.audit_log.find_one({"action": "import_code_dates"})


async def test_finance_admin_may_import(client, seeded_db):
    await seeded_db.agent_profiles.insert_one({
        "agent_id": "FA_1", "name": "Finance One", "email": "fa1@test.dev",
        "role": "finance_admin", "upline_id": None, "office": "",
    })
    token = await make_session(seeded_db, role="finance_admin", agent_id="FA_1", email="fa1@test.dev")
    r = await _upload(client, token, _workbook([("AGENT ONE", datetime(2026, 3, 2), "")]))
    assert r.status_code == 200, r.text


async def test_mga_may_not_import(client, seeded_db):
    token = await make_session(seeded_db, role="level_3", agent_id="MGA_1", email="mga1@test.dev")
    r = await _upload(client, token, _workbook([("AGENT ONE", datetime(2026, 3, 2), "")]))
    assert r.status_code == 403


async def test_sheet_without_code_date_column_is_rejected(client, seeded_db):
    token = await _admin(seeded_db)
    r = await _upload(client, token, _workbook([("AGENT ONE", "x", "")], header=("Agent", "Joined", "MGA")))
    assert r.status_code == 400


# ---------------- the Secret Sauce reads the code date, not the stale flag ----------------

async def test_secret_sauce_derives_rookie_from_code_date(client, seeded_db):
    """A flag set today goes stale on the anniversary; the sheet must judge by
    the code date as of the week it reports."""
    await seeded_db.agent_profiles.update_one(
        {"agent_id": "AG_1"}, {"$set": {"code_date": "2025-07-15", "is_rookie": True}})
    await seeded_db.production_entries.insert_one({
        "entry_id": "pe", "agent_id": "AG_1", "office": "MCM", "sales_day": "2026-07-16",
        "sets": 0, "sits": 0, "sales": 0, "ots_sits": 0, "ots_sales": 0, "n1": 0,
        "refs_obtained": 0, "ref_sits": 0, "ref_sales": 0, "pos_sits": 0, "pos_sales": 0,
        "vet_sits": 0, "vet_sales": 0, "gross_alp": 500.0, "net_alp": 500.0,
    })
    token = await make_session(seeded_db, role="level_4", agent_id="RGA_1", email="rga1@test.dev")
    # Week of 2026-07-15: one year to the day since coding — a veteran now,
    # whatever the stored flag says.
    r = await client.get("/api/vault/secret-sauce?week_start=2026-07-15", headers=auth(token))
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws["B3"].value == "Agent One", "ranked as a veteran"
    assert ws["F3"].value is None, "not in the rookie block"
